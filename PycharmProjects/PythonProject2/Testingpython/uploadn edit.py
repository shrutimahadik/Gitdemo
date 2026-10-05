from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions
from selenium.webdriver.support.select import Select
from selenium.webdriver.support.wait import WebDriverWait
import openpyxl
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions
from selenium.webdriver.support.select import Select
from selenium.webdriver.support.wait import WebDriverWait
def update_exceldata(filePath,searchTerm,colName,new_Value):
 book=openpyxl.load_workbook(filePath)
 sheet=book.active
 Dict={}

 for i in range(1,sheet.max_column+1):
     if sheet.cell(row=1, column=1).value==colName:
         Dict["col"]==i

 for i in range(1,sheet.max_row+1):
     for j in range(1, sheet.max_column + 1):
         if sheet.cell(row=1, column=j).value == searchTerm:
             Dict["row"] == i


 sheet.cell(row=Dict["row"], column=Dict["col"]).value= new_Value
 book.save(file_path)


file_path = "C:\\Users\\GNESH\\Downloads\\download.xlsx"
fruit_name="Apple"
newValue="999"
driver = webdriver.Chrome()
driver.get("https://rahulshettyacademy.com/upload-download-test/index.html")
driver.maximize_window()
driver.implicitly_wait(2)
driver.find_element(By.CSS_SELECTOR, "#downloadButton").click()

#edit the excel with update value
update_exceldata(file_path, fruit_name,"price",newValue)
file_inputexample = driver.find_element(By.CSS_SELECTOR, "#fileinput")
file_inputexample.send_keys(file_path)


wait = WebDriverWait(driver, 15)
toast_locator = (By.CSS_SELECTOR, ".Toastify__toast-body div:nth-child(2)")
wait.until(expected_conditions.presence_of_element_located(toast_locator))
print(driver.find_element(*toast_locator).text)
pricecolumn=driver.find_element(By.XPATH,"//div[text()='Price']").get_attribute("data-column-id")
#Xpath made as went to parent class-div,again parent class div from that parent class to immidiate locator
actualprice=driver.find_element(By.XPATH,"//div[text()='"+fruit_name+"']/parent::div/parent::div/div[@id='cell-"+pricecolumn+"-undefined']").text
assert actualprice==newValue