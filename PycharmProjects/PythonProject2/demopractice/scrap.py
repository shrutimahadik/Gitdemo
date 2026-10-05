import time

import openpyxl
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions
from selenium.webdriver.support.wait import WebDriverWait



filepath1="C:\\Users\\GNESH\\Downloads\\download (12).xlsx"
newaddvalue="999"
searchterm="Orange"
driver=webdriver.Chrome()
driver.implicitly_wait(10)
driver.maximize_window()
book=openpyxl.load_workbook(f)
sheet=book.active
dict={}

#filedownload
driver.get(filepath1)
driver.find_element(By.ID, "downloadButton").click()

def dataedit(filepath,serachterm,colname,newaddvalue):

    for i in range(1, sheet.max_column + 1):
        if sheet.cell(row=1,column=i).value=="price":
            dict["col"]=i

    for i in range(1,sheet.max_row+1):
        for j in range(1, sheet.max_column + 1):
            if sheet.cell(row=i,column=j).value==searchterm:
             dict["row"]=i

    sheet.cell(row=dict["row"],column=dict["col"]).value=newaddvalue
    book.save(filelocation)




#fileupload
dataedit(filelocation, searchterm, "price", newaddvalue)
path=driver.find_element(By.XPATH,"//input[@type='file']")
path.send_keys(filelocation)
wait=WebDriverWait(driver,15)
wait.until(expected_conditions.visibility_of_element_located((By.XPATH,"//div[text()='Updated Excel Data Successfully.']")))
time.sleep(3)

#valuechange
storevalue=driver.find_element(By.XPATH,"//div[@id='cell-2-undefined']/div[text()='Orange']/parent::div/parent::div/div[@data-column-id='4']").text
print(storevalue)
