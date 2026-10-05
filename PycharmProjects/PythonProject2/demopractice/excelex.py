import openpyxl

book=openpyxl.load_workbook("C:\\Users\\GNESH\\Downloads\\download.xlsx")
sheet=book.active
#cell=sheet.cell(row=1,column=2)
#print(cell.value)
#sheet.cell(row=1,column=2).value="Suru"
#print(sheet.cell(row=1,column=2).value)
#print(sheet.max_row)
#print(sheet.max_column)
#print(sheet['B5'].value)

for i in range(1,sheet.max_row+1):
    if sheet.cell(row=i,column=2).value=="Orange":
     for j in range(1,sheet.max_column+1):
      print(sheet.cell(row=i,column=j).value)


