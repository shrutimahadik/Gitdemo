import openpyxl
book=openpyxl.load_workbook("C:\\Users\\GNESH\\Downloads\\download.xlsx")
sheet=book.worksheets[1]
Dict={}
#cell=sheet.cell(row=1,column=2)
#print(cell.value)
#print(sheet.max_row)
#print(sheet.max_column)
#print(sheet['B5'].value)

for i in range (1 ,sheet.max_row+1):
    if sheet.cell(row=i,column=3).value=="mahadik":

     for j in range(1,sheet.max_column+1):
        Dict[print(sheet.cell(row=1,column=j).value)]=sheet.cell(row=i,column=j).value




